# -*- coding: utf-8 -*-
import os
import json
import threading
from pathlib import Path
from datetime import datetime
import customtkinter as ctk
from tkinter import filedialog
import docx

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

try:
    from google import genai
    from google.genai import types
    HAS_GENAI_NEW = True
except ImportError:
    HAS_GENAI_NEW = False
    try:
        import google.generativeai as legacy_genai
        HAS_GENAI_OLD = True
    except ImportError:
        HAS_GENAI_OLD = False

BASE_DIR = Path(__file__).resolve().parent
STORAGE_DIR = BASE_DIR / "storage"
OUTBOX_DIR = STORAGE_DIR / "outbox"
OUTBOX_DIR.mkdir(parents=True, exist_ok=True)

def read_text_from_file(file_path: Path) -> str:
    """Nolasa saturu no TXT, DOCX, PDF vai EXCEL faila."""
    ext = file_path.suffix.lower()
    if ext == ".txt":
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read()
        except UnicodeDecodeError:
            with open(file_path, "r", encoding="latin-1") as f:
                return f.read()
    elif ext == ".docx":
        try:
            doc = docx.Document(file_path)
            full_text = []
            for p in doc.paragraphs:
                if p.text.strip():
                    full_text.append(p.text)
            for tbl in doc.tables:
                for row in tbl.rows:
                    row_data = [c.text.strip() for c in row.cells if c.text.strip()]
                    if row_data:
                        full_text.append(" | ".join(row_data))
            return "\n".join(full_text)
        except Exception as e:
            return f"⚠️ Nevarēja nolasīt DOCX failu: {e}"
    elif ext == ".pdf":
        try:
            import pypdf
            reader = pypdf.PdfReader(str(file_path))
            pages_txt = [p.extract_text() or "" for p in reader.pages]
            return "\n".join(pages_txt)
        except ImportError:
            return "⚠️ Lai lasītu PDF, terminālī palaidiet: pip install pypdf"
        except Exception as e:
            return f"⚠️ Nevarēja nolasīt PDF: {e}"
    elif ext in [".xlsx", ".xls"]:
        try:
            import openpyxl
            wb = openpyxl.load_workbook(file_path, data_only=True)
            full_text = []
            for sheetname in wb.sheetnames:
                sheet = wb[sheetname]
                full_text.append(f"--- Lapa: {sheetname} ---")
                for row in sheet.iter_rows(values_only=True):
                    row_vals = [str(c).strip() for c in row if c is not None and str(c).strip() != "" and str(c).strip().lower() != "nan"]
                    if row_vals:
                        full_text.append(" | ".join(row_vals))
            return "\n".join(full_text)
        except Exception:
            try:
                import pandas as pd
                df = pd.read_excel(file_path)
                return df.dropna(how="all").to_string()
            except Exception as e:
                return f"⚠️ Nevarēja nolasīt Excel: {e}"
    return ""

def create_offer_excel(data: dict, margin_pct: float, lang: str = "LV") -> Path:
    """Izveido profesionālu Excel tabulu ar sarkanās zonas drošību un valodas atbalstu."""
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_file = OUTBOX_DIR / f"Piedavajums_{timestamp}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Quote" if lang == "EN" else "Piedāvājums"
    ws.views.sheetView[0].showGridLines = True

    font_header = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    font_bold = Font(name="Calibri", size=11, bold=True)
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    fill_internal_header = PatternFill(start_color="991B1B", end_color="991B1B", fill_type="solid")
    font_internal_header = Font(name="Calibri", size=11, bold=True, color="FEE2E2")
    fill_internal_data = PatternFill(start_color="FEF2F2", end_color="FEF2F2", fill_type="solid")
    border_internal = Border(
        left=Side(style='thin', color='FCA5A5'),
        right=Side(style='thin', color='FCA5A5'),
        top=Side(style='thin', color='FCA5A5'),
        bottom=Side(style='thin', color='FCA5A5')
    )

    client_display = data.get('client_name') or ("Valued Customer" if lang == "EN" else "Cienījamais Klients")
    ws.merge_cells("A1:E1")
    title_cell = ws["A1"]
    title_text = f"AuraQuadro — Commercial Offer ({client_display})" if lang == "EN" else f"AuraQuadro — Cenu Piedāvājums ({client_display})"
    title_cell.value = title_text
    title_cell.font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    title_cell.alignment = align_left

    ws.merge_cells("G1:H1")
    warn_cell = ws["G1"]
    warn_cell.value = "⚠️ INTERNAL ONLY / TIKAI IEKŠĒJAI LIETOŠANAI"
    warn_cell.font = Font(name="Calibri", size=9, bold=True, color="DC2626")
    warn_cell.alignment = align_center

    if lang == "EN":
        headers_client = ["P/N", "Model / Description", "Qty", "Unit Price EUR", "Total EUR excl. VAT"]
        headers_internal = ["Cost EUR", "Margin %"]
    else:
        headers_client = ["P/N", "Modelis / Apraksts", "Skaits", "Vienības Cena EUR", "Summa EUR bez PVN"]
        headers_internal = ["Iepirkums EUR", "Marža %"]

    ws.append([])
    header_row = 3
    for col_idx, h in enumerate(headers_client, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center if col_idx in [1, 3] else (align_left if col_idx == 2 else align_right)

    ws.cell(row=header_row, column=6, value="")

    for col_idx, h in enumerate(headers_internal, start=7):
        cell = ws.cell(row=header_row, column=col_idx, value=h)
        cell.font = font_internal_header
        cell.fill = fill_internal_header
        cell.alignment = align_center if col_idx == 8 else align_right

    items = data.get("items", [])
    start_row = 4
    for i, itm in enumerate(items, start=start_row):
        qty = int(itm.get("qty", 1) or 1)
        cost = float(itm.get("cost", 0.0) or 0.0)
        
        ws.cell(row=i, column=1, value=itm.get("pn", "")).alignment = align_center
        ws.cell(row=i, column=2, value=itm.get("desc", "")).alignment = align_left
        ws.cell(row=i, column=3, value=qty).alignment = align_center
        
        c_unit = ws.cell(row=i, column=4, value=f"=ROUND(G{i}*(1+H{i}/100), 2)")
        c_unit.alignment = align_right
        c_unit.number_format = '#,##0.00'

        c_tot = ws.cell(row=i, column=5, value=f"=ROUND(C{i}*D{i}, 2)")
        c_tot.alignment = align_right
        c_tot.number_format = '#,##0.00'

        for c in range(1, 6):
            ws.cell(row=i, column=c).border = thin_border

        c_cost = ws.cell(row=i, column=7, value=cost)
        c_cost.alignment = align_right
        c_cost.number_format = '#,##0.00'
        c_cost.fill = fill_internal_data
        c_cost.border = border_internal

        c_mrg = ws.cell(row=i, column=8, value=margin_pct)
        c_mrg.alignment = align_center
        c_mrg.number_format = '0.0"%"'
        c_mrg.fill = fill_internal_data
        c_mrg.border = border_internal

    end_row = start_row + len(items) - 1 if items else start_row

    r_sum = end_row + 2
    lbl_sub = "Total excl. VAT:" if lang == "EN" else "Kopā bez PVN:"
    lbl_vat = "VAT 21%:" if lang == "EN" else "PVN 21%:"
    lbl_tot = "TOTAL INCL. VAT:" if lang == "EN" else "KOPĀ AR PVN:"

    ws[f"D{r_sum}"] = lbl_sub
    ws[f"D{r_sum}"].font = font_bold
    ws[f"D{r_sum}"].alignment = align_right
    ws[f"E{r_sum}"] = f"=SUM(E{start_row}:E{end_row})"
    ws[f"E{r_sum}"].font = font_bold
    ws[f"E{r_sum}"].alignment = align_right
    ws[f"E{r_sum}"].number_format = '#,##0.00'

    ws[f"D{r_sum+1}"] = lbl_vat
    ws[f"D{r_sum+1}"].font = font_bold
    ws[f"D{r_sum+1}"].alignment = align_right
    ws[f"E{r_sum+1}"] = f"=ROUND(E{r_sum}*0.21, 2)"
    ws[f"E{r_sum+1}"].font = font_bold
    ws[f"E{r_sum+1}"].alignment = align_right
    ws[f"E{r_sum+1}"].number_format = '#,##0.00'

    ws[f"D{r_sum+2}"] = lbl_tot
    ws[f"D{r_sum+2}"].font = Font(name="Calibri", size=12, bold=True, color="059669")
    ws[f"D{r_sum+2}"].alignment = align_right
    ws[f"E{r_sum+2}"] = f"=E{r_sum}+E{r_sum+1}"
    ws[f"E{r_sum+2}"].font = Font(name="Calibri", size=12, bold=True, color="059669")
    ws[f"E{r_sum+2}"].alignment = align_right
    ws[f"E{r_sum+2}"].number_format = '#,##0.00'

    col_widths = {"A": 14, "B": 48, "C": 10, "D": 18, "E": 20, "F": 4, "G": 16, "H": 12}
    for col, width in col_widths.items():
        ws.column_dimensions[col].width = width

    wb.save(out_file)
    return out_file

def generate_client_offer_structured(client_req: str, dist_costs: str, margin_pct: float, target_lang: str, api_key: str):
    """Gemini analīze: izvelk datus, skaidrojumus licencēm/servisam un formē tekstu izvēlētajā valodā."""
    if not api_key:
        return {"error": "Trūkst Gemini API atslēgas"}

    lang_instruction = {
        "LV": "Valoda: LATVIEŠU. Sagatavo tekstu un paskaidrojumus profesionālā latviešu valodā.",
        "EN": "Language: ENGLISH. Prepare all descriptions, notes, and greetings in professional business English.",
        "Auto": "Valoda: Noteikt automātiski pēc klienta pieprasījuma (ja latviski - LV, ja angliski - EN). Noklusējums: LV."
    }.get(target_lang, "Valoda: LATVIEŠU.")

    prompt = f"""Tu esi tirdzniecības un tāmēšanas inženieris.
Izanalizē datus un atgriez strukturētu informāciju JSON formātā.

{lang_instruction}

--- 1. KLIENTA PIEPRASĪJUMS ---
{client_req if client_req.strip() else "(Nav atsevišķi norādīts, izmanto piegādātāja specifikāciju)"}

--- 2. PIEGĀDĀTĀJA CENAS UN POZĪCIJAS ---
{dist_costs}

STINGRI NOTEIKUMI:
1. Laukā "cost" norādi TIKAI oriģinālo piegādātāja vienības pašizmaksu (skaitli) bez maržas.
2. Modeļa aprakstā ("desc") iekļauj faktuālu aprakstu no ražotāja specifikācijas/datasheet. Nefantazē.
3. LICENCES / SERVISA PAKAS:
   - Saraksta beigās iekļauj programmatūras licences, Aruba Central, Care Pack vai servisa pakas.
   - Laukā "license_notes" sniedz skaidru un lietderīgu aprakstu, kāpēc šīs licences/servisa pakas ir nepieciešamas (piem., mākoņvadība, monitorings, garantijas paplašinājums).
4. Laukā "detected_lang" norādi valodas kodu ("LV" vai "EN").

Atgriez TIKAI derīgu JSON šādā struktūrā:
{{
  "detected_lang": "LV",
  "client_name": "Klienta nosaukums vai uzņēmums",
  "delivery_notes": "Piegādes termiņi un garantijas noteikumi",
  "license_notes": "Skaidrojums par licencēm un servisa pakām un to sniegtajām priekšrocībām",
  "items": [
    {{
      "pn": "Artikuls vai kods",
      "desc": "Modeļa nosaukums un apraksts",
      "qty": 1,
      "cost": 100.00
    }}
  ]
}}
"""

    try:
        raw = ""
        if HAS_GENAI_NEW:
            client = genai.Client(api_key=api_key)
            response = client.models.generate_content(
                model="gemini-2.5-flash",
                contents=prompt,
                config=types.GenerateContentConfig(response_mime_type="application/json", temperature=0.1)
            )
            raw = response.text.strip()
        elif HAS_GENAI_OLD:
            legacy_genai.configure(api_key=api_key)
            model = legacy_genai.GenerativeModel(
                model_name="gemini-2.5-flash",
                generation_config={"response_mime_type": "application/json", "temperature": 0.1}
            )
            res = model.generate_content(prompt)
            raw = res.text.strip()

        data = json.loads(raw)
        effective_lang = target_lang if target_lang in ["LV", "EN"] else data.get("detected_lang", "LV").upper()
        data["final_lang"] = effective_lang

        # E-pasta teksta veidošana atbilstoši valodai
        client_name = data.get("client_name")
        if effective_lang == "EN":
            greeting = f"Dear {client_name if client_name else 'Customer'},"
            intro = "Thank you for your inquiry. Please find our commercial quotation below:"
            col_pn, col_desc, col_qty, col_price, col_tot = "P/N", "Model / Description", "Qty", "Price excl.", "Total EUR"
            lbl_sub, lbl_vat, lbl_tot = "Total excl. VAT:", "VAT 21%:", "TOTAL INCL. VAT:"
            lbl_delivery, lbl_licenses = "Delivery terms:", "Software / Service Packages:"
            signoff = "Sincerely,\nAuraQuadro"
        else:
            greeting = f"Labdien{', ' + client_name if client_name else ''}!"
            intro = "Pateicamies par pieprasījumu. Nosūtām sagatavoto cenu piedāvājumu:"
            col_pn, col_desc, col_qty, col_price, col_tot = "P/N", "Modelis / Apraksts", "Skaits", "Cena bez PVN", "Kopā EUR"
            lbl_sub, lbl_vat, lbl_tot = "Kopā bez PVN:", "PVN 21%:", "KOPĀ AR PVN:"
            lbl_delivery, lbl_licenses = "Piegādes nosacījumi:", "Licences un servisa pakas:"
            signoff = "Ar cieņu,\nAuraQuadro"

        lines = [
            greeting + "\n",
            intro + "\n",
            "-" * 72,
            f"{col_pn:<12} | {col_desc:<30} | {col_qty:<6} | {col_price:<12} | {col_tot}",
            "-" * 72
        ]
        
        tot_no_pvn = 0.0
        for itm in data.get("items", []):
            qty = int(itm.get("qty", 1) or 1)
            cost = float(itm.get("cost", 0.0) or 0.0)
            unit_price = round(cost * (1 + margin_pct / 100), 2)
            row_tot = round(qty * unit_price, 2)
            tot_no_pvn += row_tot
            
            pn = str(itm.get("pn", ""))[:12]
            desc = str(itm.get("desc", ""))[:30]
            lines.append(f"{pn:<12} | {desc:<30} | {qty:<6} | {unit_price:>12.2f} | {row_tot:>10.2f}")

        pvn = round(tot_no_pvn * 0.21, 2)
        tot_with_pvn = round(tot_no_pvn + pvn, 2)

        lines.extend([
            "-" * 72,
            f"{lbl_sub:<16} {tot_no_pvn:>10.2f} EUR",
            f"{lbl_vat:<16} {pvn:>10.2f} EUR",
            f"{lbl_tot:<16} {tot_with_pvn:>10.2f} EUR\n"
        ])

        # Piezīmes par licencēm un piegādi
        if data.get("license_notes"):
            lines.append(f"💡 {lbl_licenses}\n{data.get('license_notes')}\n")
        
        delivery_txt = data.get("delivery_notes") or ("Delivery according to agreement. Manufacturer warranty." if effective_lang == "EN" else "Piegādes termiņš pēc vienošanās. Ražotāja garantija.")
        lines.append(f"📦 {lbl_delivery} {delivery_txt}\n")
        lines.append(signoff)

        data["email_text"] = "\n".join(lines)
        return data

    except Exception as e:
        return {"error": str(e), "email_text": f"⚠️ Kļūda aprēķinā: {e}", "items": [], "final_lang": "LV"}

class OfferStudioFrame(ctk.CTkFrame):
    def __init__(self, master, get_gemini_key_fn=None):
        super().__init__(master, fg_color="transparent")
        self.get_gemini_key_fn = get_gemini_key_fn
        self.last_generated_excel = None

        # Vadības josla
        mode_bar = ctk.CTkFrame(self, fg_color="#0f172a", height=50)
        mode_bar.pack(fill="x", padx=10, pady=(10, 5))

        # Maržas ievade
        ctk.CTkLabel(mode_bar, text="Marža (%):", font=ctk.CTkFont(size=12, weight="bold"), text_color="#38bdf8").pack(side="left", padx=(15, 5))
        self.entry_margin = ctk.CTkEntry(mode_bar, width=50, height=30, font=ctk.CTkFont(size=12, weight="bold"))
        self.entry_margin.insert(0, "12")
        self.entry_margin.pack(side="left", padx=(0, 10))

        # Valodas selektors
        ctk.CTkLabel(mode_bar, text="Valoda:", font=ctk.CTkFont(size=12, weight="bold"), text_color="#cbd5e1").pack(side="left", padx=(5, 5))
        self.seg_lang = ctk.CTkSegmentedButton(
            mode_bar,
            values=["Auto", "LV", "EN"],
            width=130,
            height=30
        )
        self.seg_lang.set("Auto")
        self.seg_lang.pack(side="left", padx=(0, 12))

        self.btn_calculate = ctk.CTkButton(
            mode_bar,
            text="⚡ Aprēķināt & Noformēt",
            width=180,
            height=32,
            fg_color="#2563eb",
            hover_color="#1d4ed8",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.process_offer
        )
        self.btn_calculate.pack(side="left", padx=5, pady=8)

        self.btn_open_excel = ctk.CTkButton(
            mode_bar,
            text="📊 Atvērt Excel Tāmēšanai",
            width=180,
            height=32,
            fg_color="#0284c7",
            hover_color="#0369a1",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.open_generated_excel,
            state="disabled"
        )
        self.btn_open_excel.pack(side="left", padx=5, pady=8)

        self.btn_copy = ctk.CTkButton(
            mode_bar,
            text="📋 Kopēt E-pastam",
            width=130,
            height=32,
            fg_color="#059669",
            hover_color="#047857",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.copy_to_clipboard
        )
        self.btn_copy.pack(side="left", padx=5, pady=8)

        self.lbl_status = ctk.CTkLabel(mode_bar, text="", font=ctk.CTkFont(size=12), text_color="#38bdf8")
        self.lbl_status.pack(side="right", padx=15)

        # Sadalījums 2 kolonnās
        panes = ctk.CTkFrame(self, fg_color="transparent")
        panes.pack(fill="both", expand=True, padx=10, pady=5)

        # Kreisā kolonna
        left_col = ctk.CTkFrame(panes, fg_color="#0b1329", corner_radius=8)
        left_col.pack(side="left", fill="both", expand=True, padx=(0, 5))

        h1 = ctk.CTkFrame(left_col, fg_color="transparent")
        h1.pack(fill="x", padx=10, pady=(6, 2))
        ctk.CTkLabel(h1, text="1. 📥 Klienta Pieprasījums:", font=ctk.CTkFont(size=12, weight="bold"), text_color="#94a3b8").pack(side="left")
        ctk.CTkButton(
            h1, text="📁 Ielādēt", width=70, height=22, fg_color="#334155", hover_color="#475569", font=ctk.CTkFont(size=10),
            command=lambda: self.load_into_textbox(self.txt_client)
        ).pack(side="right")

        self.txt_client = ctk.CTkTextbox(left_col, fg_color="#0f172a", font=ctk.CTkFont(size=11), wrap="word", height=140)
        self.txt_client.pack(fill="x", padx=10, pady=(0, 8))

        h2 = ctk.CTkFrame(left_col, fg_color="transparent")
        h2.pack(fill="x", padx=10, pady=(4, 2))
        ctk.CTkLabel(h2, text="2. 🏷️ Distributora Pašizmaksa / Cenas:", font=ctk.CTkFont(size=12, weight="bold"), text_color="#38bdf8").pack(side="left")
        ctk.CTkButton(
            h2, text="📁 Ielādēt (.xlsx)", width=105, height=22, fg_color="#334155", hover_color="#475569", font=ctk.CTkFont(size=10),
            command=lambda: self.load_into_textbox(self.txt_dist)
        ).pack(side="right")

        self.txt_dist = ctk.CTkTextbox(left_col, fg_color="#0f172a", font=ctk.CTkFont(size=11), wrap="word")
        self.txt_dist.pack(fill="both", expand=True, padx=10, pady=(0, 8))

        # Labā kolonna
        right_col = ctk.CTkFrame(panes, fg_color="#0b1329", corner_radius=8)
        right_col.pack(side="right", fill="both", expand=True, padx=(5, 0))

        right_header = ctk.CTkFrame(right_col, fg_color="transparent")
        right_header.pack(fill="x", padx=10, pady=(6, 2))
        ctk.CTkLabel(right_header, text="⚡ Gatavais Piedāvājums Klientam (Drošs):", font=ctk.CTkFont(size=12, weight="bold"), text_color="#10b981").pack(side="left")

        self.txt_preview = ctk.CTkTextbox(right_col, fg_color="#0f172a", font=ctk.CTkFont(size=12), wrap="word")
        self.txt_preview.pack(fill="both", expand=True, padx=10, pady=(0, 8))

    def load_into_textbox(self, target_widget):
        f = filedialog.askopenfilename(
            filetypes=[("Dokumenti & Tabulas", "*.xlsx *.xls *.docx *.pdf *.txt"), ("Visi faili", "*.*")]
        )
        if f:
            path = Path(f)
            self.lbl_status.configure(text=f"Ielādē: {path.name}...")
            content = read_text_from_file(path)
            if content:
                target_widget.delete("1.0", "end")
                target_widget.insert("end", content)
                self.lbl_status.configure(text=f"Ielādēts: {path.name}")
            else:
                self.lbl_status.configure(text="⚠️ Neizdevās nolasīt tekstu")

    def load_request(self, subject, sender, body):
        self.txt_client.delete("1.0", "end")
        self.txt_client.insert("end", f"Temats: {subject}\nNo: {sender}\n\n{body}")
        self.lbl_status.configure(text="Pieprasījums ielādēts no Radara!")

    def process_offer(self):
        client_text = self.txt_client.get("1.0", "end").strip()
        dist_text = self.txt_dist.get("1.0", "end").strip()

        if not dist_text:
            self.lbl_status.configure(text="⚠️ Ievadiet distributora cenas!")
            return

        try:
            margin = float(self.entry_margin.get().replace("%", "").strip())
        except ValueError:
            margin = 12.0

        target_lang = self.seg_lang.get()
        api_key = self.get_gemini_key_fn() if self.get_gemini_key_fn else os.environ.get("GEMINI_API_KEY", "")
        self.lbl_status.configure(text="Gemini tāmē & ģenerē Excel...")
        self.btn_calculate.configure(state="disabled")

        def _worker():
            excel_path = None
            email_out = ""
            status_msg = "Gatavs!"
            try:
                data = generate_client_offer_structured(client_text, dist_text, margin, target_lang, api_key)
                email_out = data.get("email_text", "")
                final_lang = data.get("final_lang", "LV")
                if HAS_OPENPYXL and data.get("items"):
                    try:
                        excel_path = create_offer_excel(data, margin, lang=final_lang)
                        status_msg = f"Gatavs! Saglabāts {excel_path.name}"
                    except Exception as e:
                        email_out += f"\n\n[Piezīme: Neizdevās saglabāt Excel: {e}]"
            except Exception as e:
                email_out = f"⚠️ Kļūda: {e}"
                status_msg = "Kļūda!"
            finally:
                def _done():
                    self.txt_preview.delete("1.0", "end")
                    self.txt_preview.insert("end", email_out)
                    self.btn_calculate.configure(state="normal")
                    if excel_path and excel_path.exists():
                        self.last_generated_excel = excel_path
                        self.btn_open_excel.configure(state="normal")
                    self.lbl_status.configure(text=status_msg)
                self.after(0, _done)

        threading.Thread(target=_worker, daemon=True).start()

    def open_generated_excel(self):
        if self.last_generated_excel and self.last_generated_excel.exists():
            try:
                os.startfile(str(self.last_generated_excel))
                self.lbl_status.configure(text="Excel atvērts tāmēšanai!")
            except Exception as e:
                self.lbl_status.configure(text=f"Kļūda atverot Excel: {e}")

    def copy_to_clipboard(self):
        txt = self.txt_preview.get("1.0", "end").strip()
        if txt:
            self.clipboard_clear()
            self.clipboard_append(txt)
            self.lbl_status.configure(text="📋 E-pasta teksts iekopēts!")

    def set_ui_language(self, lang: str):
        """Dinamiski atjauno visus OfferStudio tekstus atkarībā no izvēlētās valodas."""
        self.seg_lang.set(lang)
        if lang == "EN":
            self.btn_calculate.configure(text="⚡ Calculate & Generate")
            self.btn_open_excel.configure(text="📊 Open Excel Sheet")
            self.btn_copy.configure(text="📋 Copy for Email")
            self.lbl_margin_text.configure(text="Margin (%):")
            self.lbl_lang_text.configure(text="Language:")
            self.lbl_client_header.configure(text="1. 📥 Client RFQ / Requirements:")
            self.btn_load_client.configure(text="📁 Load")
            self.lbl_dist_header.configure(text="2. 🏷️ Supplier / Disti Costs:")
            self.btn_load_dist.configure(text="📁 Load (.xlsx)")
            self.lbl_result_header.configure(text="⚡ Commercial Offer (Client-Safe):")
        else:
            self.btn_calculate.configure(text="⚡ Aprēķināt & Noformēt")
            self.btn_open_excel.configure(text="📊 Atvērt Excel Tāmēšanai")
            self.btn_copy.configure(text="📋 Kopēt E-pastam")
            self.lbl_margin_text.configure(text="Marža (%):")
            self.lbl_lang_text.configure(text="Valoda:")
            self.lbl_client_header.configure(text="1. 📥 Klienta Pieprasījums:")
            self.btn_load_client.configure(text="📁 Ielādēt")
            self.lbl_dist_header.configure(text="2. 🏷️ Distributora Pašizmaksa / Cenas:")
            self.btn_load_dist.configure(text="📁 Ielādēt (.xlsx)")
            self.lbl_result_header.configure(text="⚡ Gatavais Piedāvājums Klientam (Drošs):")
