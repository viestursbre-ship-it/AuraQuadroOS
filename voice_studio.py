# -*- coding: utf-8 -*-
import os
import base64
import threading
from tkinter import filedialog
import customtkinter as ctk
import google.generativeai as genai

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False

class VoiceStudioFrame(ctk.CTkFrame):
    def __init__(self, master, get_gemini_key_fn=None, **kwargs):
        super().__init__(master, fg_color="transparent", **kwargs)
        self.get_gemini_key_fn = get_gemini_key_fn
        self.selected_audio_path = None
        self._build_ui()

    def _build_ui(self):
        # Augšējā josla faila izvēlei
        top_bar = ctk.CTkFrame(self, fg_color="#0f172a")
        top_bar.pack(fill="x", padx=10, pady=10)

        self.btn_select_audio = ctk.CTkButton(
            top_bar, 
            text="📁 Izvēlēties audio (.ogg, .mp3, .m4a, .wav)", 
            fg_color="#2563eb", 
            command=self.choose_audio_file
        )
        self.btn_select_audio.pack(side="left", padx=10, pady=10)

        self.lbl_audio_name = ctk.CTkLabel(top_bar, text="Nav izvēlēts neviens audio fails", text_color="#94a3b8")
        self.lbl_audio_name.pack(side="left", padx=10)

        # Režīmu slēdži labajā pusē
        self.voice_mode = ctk.StringVar(value="rezume")
        rb1 = ctk.CTkRadioButton(top_bar, text="Klienta Rezumē", variable=self.voice_mode, value="rezume")
        rb1.pack(side="right", padx=15)
        rb2 = ctk.CTkRadioButton(top_bar, text="Sēdes Protokols", variable=self.voice_mode, value="protokols")
        rb2.pack(side="right", padx=10)

        # Galvenais teksta lauks
        self.txt_voice_output = ctk.CTkTextbox(
            self, 
            font=ctk.CTkFont(size=13), 
            fg_color="#0f172a", 
            text_color="#f1f5f9", 
            wrap="word"
        )
        self.txt_voice_output.pack(fill="both", expand=True, padx=10, pady=5)
        self.txt_voice_output.insert("end", "🎙️ Balss Studija gatava darbam.\nIzvēlieties WhatsApp .ogg failu un spiediet 'Apstrādāt ar Gemini MI'.\n\n")

        # Apakšējā pogu josla
        bottom_bar = ctk.CTkFrame(self, fg_color="transparent")
        bottom_bar.pack(fill="x", padx=10, pady=10)

        self.btn_run_voice = ctk.CTkButton(
            bottom_bar, 
            text="⚡ Apstrādāt ar Gemini MI", 
            height=40, 
            fg_color="#059669", 
            hover_color="#10b981", 
            command=self.process_audio
        )
        self.btn_run_voice.pack(side="left", fill="x", expand=True, padx=(0, 10))

        self.btn_export_docx = ctk.CTkButton(
            bottom_bar, 
            text="💾 Saglabāt Word (.docx)", 
            height=40, 
            width=220,
            fg_color="#3b82f6", 
            command=self.export_to_docx
        )
        self.btn_export_docx.pack(side="right", padx=(10, 0))

    def choose_audio_file(self):
        f = filedialog.askopenfilename(
            filetypes=[("Audio faili", "*.ogg *.mp3 *.m4a *.wav"), ("Visi faili", "*.*")]
        )
        if f:
            self.selected_audio_path = f
            self.lbl_audio_name.configure(text=os.path.basename(f), text_color="#38bdf8")

    def process_audio(self):
        if not self.selected_audio_path:
            self.txt_voice_output.insert("end", "⚠️ Lūdzu, vispirms izvēlieties audio failu!\n")
            return

        api_key = self.get_gemini_key_fn() if self.get_gemini_key_fn else os.environ.get("GEMINI_API_KEY", "")
        if not api_key:
            self.txt_voice_output.insert("end", "⚠️ Gemini API atslēga nav atrasta!\n")
            return

        self.txt_voice_output.delete("1.0", "end")
        self.txt_voice_output.insert("end", "⏳ Gemini klausās balss ziņu un ģenerē struktūru...\n")
        self.btn_run_voice.configure(state="disabled")

        threading.Thread(target=self._run_audio_ai, args=(api_key,), daemon=True).start()

    def _run_audio_ai(self, api_key):
        out_text = ""
        try:
            genai.configure(api_key=api_key)
            with open(self.selected_audio_path, "rb") as f:
                audio_bytes = f.read()

            b64_data = base64.b64encode(audio_bytes).decode("utf-8")
            ext = os.path.splitext(self.selected_audio_path)[1].lower()
            mime = "audio/ogg" if ext == ".ogg" else ("audio/mp3" if ext == ".mp3" else "audio/wav")

            mode = self.voice_mode.get()
            if mode == "protokols":
                prompt_text = (
                    "Tu esi precīzs lietvedis un sēžu protokolētājs. Noklausies šo audio ierakstu un sagatavo sēdes protokolu latviešu valodā:\n\n"
                    "1. SANĀKSMES TEMATS UN DALĪBNIEKI\n"
                    "2. IZSKATĪTIE JAUTĀJUMI (hronoloģiski)\n"
                    "3. PIEŅEMTIE LĒMUMI\n"
                    "4. UZDEVUMI UN TERMIŅI (kas kam jādara)\n"
                    "5. PILNA TRANSSKRIPCIJA (būtiskākais runātais teksts)."
                )
            else:
                prompt_text = (
                    "Tu esi augstākā līmeņa biznesa analītiķis. Klients ir iesūtījis haotisku balss ziņu (WhatsApp audio).\n"
                    "Pilnībā nofiltrē visas emocijas, liekvārdību, pauzes un 'bla-bla'. Izveido tīru un strukturētu kopsavilkumu latviešu valodā:\n\n"
                    "### 📋 KLIENTA PIEPRASĪJUMA REZUMĒ\n"
                    "* **Klients / Iesaistītie:** [Vārds, kontakti vai organizācija, ja minēts]\n"
                    "* **Galvenā vajadzība / Prasība:** [Sausais atlikums — kas tieši tiek prasīts]\n"
                    "* **Tehniskās nianses / Daudzumi:** [Iekārtas, apjomi, licences, parametri]\n"
                    "* **Termiņš / Steidzamība:** [Līdz kuram laikam vajag piedāvājumu vai izpildi]\n\n"
                    "### 📝 VĀRDS VĀRDĀ TRANSSKRIPCIJA\n"
                    "[Tīrs latviskots runas teksts bez starpsaucieniem]."
                )

            model = genai.GenerativeModel("gemini-2.5-flash")
            res = model.generate_content([
                {"mime_type": mime, "data": b64_data},
                prompt_text
            ])
            out_text = res.text
        except Exception as e:
            out_text = f"⚠️ Kļūda audio apstrādē: {e}"

        def update():
            self.txt_voice_output.delete("1.0", "end")
            self.txt_voice_output.insert("end", out_text)
            self.btn_run_voice.configure(state="normal")

        self.after(0, update)

    def export_to_docx(self):
        text = self.txt_voice_output.get("1.0", "end").strip()
        if not text:
            return

        if not HAS_DOCX:
            self.txt_voice_output.insert("end", "\n⚠️ python-docx bibliotēka nav instalēta! Palaidiet: pip install python-docx")
            return

        save_path = filedialog.asksaveasfilename(
            defaultextension=".docx",
            filetypes=[("Word Dokumenti", "*.docx")],
            initialfile="AQ_Balss_Rezume.docx"
        )
        if save_path:
            doc = docx.Document()
            doc.add_heading("AQ-OS Balss Apstrādes Rezultāts", level=1)
            for line in text.split("\n"):
                if line.startswith("### "):
                    doc.add_heading(line.replace("### ", ""), level=2)
                elif line.startswith("* "):
                    doc.add_paragraph(line.replace("* ", ""), style='List Bullet')
                else:
                    if line.strip():
                        doc.add_paragraph(line)

            doc.save(save_path)
            self.txt_voice_output.insert("end", f"\n\n✅ Dokuments veiksmīgi saglabāts: {save_path}")
