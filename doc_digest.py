# -*- coding: utf-8 -*-
import os
import json
import shutil
import threading
from pathlib import Path
from datetime import datetime
import customtkinter as ctk
from tkinter import filedialog

import docx
from pptx import Presentation
import openpyxl
from openpyxl.styles import Font
import google.generativeai as genai

BASE_DIR = Path(__file__).resolve().parent
STORAGE_DIR = BASE_DIR / "storage"
INBOX_DIR = STORAGE_DIR / "inbox"
OUTBOX_DIR = STORAGE_DIR / "outbox"
ARCHIVE_DIR = STORAGE_DIR / "archive"
EXCEL_REGISTRY_PATH = STORAGE_DIR / "Ligumu_Registrs.xlsx"

for d in [INBOX_DIR, OUTBOX_DIR, ARCHIVE_DIR]:
    d.mkdir(parents=True, exist_ok=True)

# ==========================================
# PALĪGFUNKCIJAS UN EXCEL REĢISTRS
# ==========================================
def safe_archive_file(source_path: Path, target_dir: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    destination = target_dir / source_path.name
    if destination.exists():
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        destination = target_dir / f"{source_path.stem}_{timestamp}{source_path.suffix}"
    shutil.move(str(source_path), str(destination))
    return destination

def append_to_excel_registry(file_name: str, analysis: dict, status: str = "Apstrādāts"):
    headers = ["Ieraksta Datums", "Fails", "Dokumenta tips", "Dokumenta Nr.", 
               "Dokumenta Datums", "Puses", "Summa", "Termiņš", "Statuss"]
    try:
        if not EXCEL_REGISTRY_PATH.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Līgumu Reģistrs"
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            wb.save(EXCEL_REGISTRY_PATH)

        wb = openpyxl.load_workbook(EXCEL_REGISTRY_PATH)
        ws = wb.active

        parties_data = analysis.get("parties", [])
        if isinstance(parties_data, list):
            parties_str = ", ".join([p.get("name", "").strip() if isinstance(p, dict) else str(p) for p in parties_data])
        else:
            parties_str = str(parties_data or "N/A")

        fin = analysis.get("financial_terms", {})
        sum_val = fin.get("total_amount", "N/A") if isinstance(fin, dict) else "N/A"
        term_val = fin.get("deadlines", "N/A") if isinstance(fin, dict) else "N/A"

        row = [
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            file_name,
            str(analysis.get("document_type", "N/A")),
            str(analysis.get("document_number", "N/A")),
            str(analysis.get("document_date", "N/A")),
            parties_str,
            str(sum_val),
            str(term_val),
            status
        ]
        ws.append(row)
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
        wb.save(EXCEL_REGISTRY_PATH)
    except Exception as e:
        print(f"⚠️ Kļūda Excel reģistrācijā: {e}")

# ==========================================
# DOKUMENTU APSTRĀDES PAKALPOJUMS
# ==========================================
class DocDigestService:
    def __init__(self, get_gemini_key_fn=None):
        self.get_gemini_key_fn = get_gemini_key_fn

    def _get_key(self):
        if self.get_gemini_key_fn:
            k = self.get_gemini_key_fn()
            if k:
                return k
        return os.environ.get("GEMINI_API_KEY") or os.environ.get("AQ_AI_API_KEY", "")

    def _read_docx(self, path: Path) -> str:
        doc = docx.Document(path)
        full_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    full_text.append(" | ".join(cells))
        return "\n".join(full_text)

    def _read_txt(self, path: Path) -> str:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def _read_pptx(self, path: Path) -> str:
        prs = Presentation(path)
        full_text = []
        for i, slide in enumerate(prs.slides, 1):
            full_text.append(f"\n--- Slaids {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            full_text.append(paragraph.text.strip())
        return "\n".join(full_text)

    def _analyze_with_gemini(self, content_payload, file_type: str = "docx") -> dict:
        api_key = self._get_key()
        if not api_key:
            return {"error": "Nav norādīta Gemini API atslēga."}

        genai.configure(api_key=api_key)
        prompt_text = """Analizējiet šo dokumentu un sagatavojiet precīzu JSON objektu:
{
  "document_type": "Līgums / Rēķins / Akts / Specifikācija / Cits",
  "document_number": "Numurs vai N/A",
  "document_date": "Datums vai N/A",
  "parties": [{"name": "Puses nosaukums", "role": "Pasūtītājs/Piegādātājs"}],
  "subject": "Darījuma priekšmets",
  "financial_terms": {"total_amount": "Summa ar PVN", "deadlines": "Apmaksas/Izpildes termiņi"},
  "risks_warnings": ["Konstatētie riski vai brīdinājumi"]
}
Atgrieziet TIKAI derīgu JSON bez markdown blokiem."""

        try:
            model = genai.GenerativeModel(
                model_name="gemini-2.5-flash",
                generation_config={"response_mime_type": "application/json", "temperature": 0.2}
            )
            
            if file_type == "pdf":
                uploaded = genai.upload_file(str(content_payload))
                response = model.generate_content([prompt_text, uploaded])
            else:
                response = model.generate_content(f"{prompt_text}\n\nDokumenta saturs:\n{content_payload[:20000]}")

            if response and response.text:
                return json.loads(response.text.strip())
            return {"error": "Tukša atbilde no MI"}
        except Exception as e:
            return {"error": str(e)}

    def _create_digest_docx(self, output_path: Path, data: dict, original_filename: str):
        doc = docx.Document()
        doc.add_heading(f"Kopsavilkums: {original_filename}", level=1)
        doc.add_paragraph(f"Izveidots: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        def add_sec(title, val):
            doc.add_heading(title, level=2)
            if isinstance(val, dict):
                for k, v in val.items():
                    doc.add_paragraph(f"• {k.replace('_', ' ').capitalize()}: {v}")
            elif isinstance(val, list):
                for item in val:
                    doc.add_paragraph(f"• {item}")
            else:
                doc.add_paragraph(str(val or "N/A"))

        add_sec("Vispārīgā informācija", {
            "Tips": data.get("document_type", "N/A"),
            "Numurs": data.get("document_number", "N/A"),
            "Datums": data.get("document_date", "N/A"),
            "Priekšmets": data.get("subject", "N/A")
        })
        add_sec("Finanšu noteikumi", data.get("financial_terms", {}))
        add_sec("Riski un brīdinājumi", data.get("risks_warnings", []))
        doc.save(output_path)

    def process_file(self, file_path: Path):
        ext = file_path.suffix.lower()
        if ext == ".docx":
            payload = self._read_docx(file_path)
        elif ext == ".txt":
            payload = self._read_txt(file_path)
        elif ext == ".pptx":
            payload = self._read_pptx(file_path)
        elif ext == ".pdf":
            payload = file_path
        else:
            return None, "Neatbalstīts formāts"

        if not payload:
            return None, "Neizdevās nolasīt faila saturu"

        analysis = self._analyze_with_gemini(payload, file_type=ext.replace(".", ""))
        if "error" in analysis:
            return None, analysis["error"]

        out_name = f"DIGEST_{file_path.stem}.docx"
        out_path = OUTBOX_DIR / out_name
        self._create_digest_docx(out_path, analysis, file_path.name)
        append_to_excel_registry(file_path.name, analysis)
        safe_archive_file(file_path, ARCHIVE_DIR)
        return analysis, str(out_path)

# ==========================================
# DOKUMENTU CILNE (CTkFrame)
# ==========================================
class DocDigestFrame(ctk.CTkFrame):
    def __init__(self, master, get_gemini_key_fn=None):
        super().__init__(master, fg_color="transparent")
        self.service = DocDigestService(get_gemini_key_fn=get_gemini_key_fn)

        top_bar = ctk.CTkFrame(self, fg_color="#0f172a", height=50)
        top_bar.pack(fill="x", padx=10, pady=(10, 5))

        self.btn_select = ctk.CTkButton(
            top_bar,
            text="📥 Izvēlēties Dokumentu",
            width=170,
            height=34,
            fg_color="#2563eb",
            hover_color="#1d4ed8",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.select_and_process
        )
        self.btn_select.pack(side="left", padx=12, pady=8)

        self.btn_open_excel = ctk.CTkButton(
            top_bar,
            text="📊 Atvērt Līgumu Reģistru",
            width=170,
            height=34,
            fg_color="#059669",
            hover_color="#10b981",
            font=ctk.CTkFont(size=12, weight="bold"),
            command=self.open_registry
        )
        self.btn_open_excel.pack(side="left", padx=6, pady=8)

        self.lbl_status = ctk.CTkLabel(top_bar, text="", font=ctk.CTkFont(size=12), text_color="#38bdf8")
        self.lbl_status.pack(side="right", padx=15)

        self.txt_output = ctk.CTkTextbox(
            self,
            font=ctk.CTkFont(size=13),
            fg_color="#0f172a",
            text_color="#f1f5f9",
            wrap="word"
        )
        self.txt_output.pack(fill="both", expand=True, padx=10, pady=10)
        self.txt_output.insert("end", "📄 Dokumentu Drop-Zone gatava.\nIelādējiet PDF, DOCX, PPTX vai TXT failu analīzei, Word kopsavilkuma izveidei un Excel reģistrēšanai.\n\n")

    def open_registry(self):
        if EXCEL_REGISTRY_PATH.exists():
            os.startfile(str(EXCEL_REGISTRY_PATH))
        else:
            self.lbl_status.configure(text="Reģistrs vēl nav izveidots")

    def select_and_process(self):
        path = filedialog.askopenfilename(
            filetypes=[("Dokumenti", "*.pdf *.docx *.pptx *.txt"), ("Visi faili", "*.*")]
        )
        if not path:
            return

        p = Path(path)
        self.btn_select.configure(state="disabled")
        self.lbl_status.configure(text=f"Apstrādā {p.name}...")

        def _worker():
            analysis, res_msg = self.service.process_file(p)
            def _done():
                self.btn_select.configure(state="normal")
                if analysis:
                    self.lbl_status.configure(text="✅ Apstrāde pabeigta!")
                    self.txt_output.insert("end", f"\n=== {p.name} ===\n")
                    self.txt_output.insert("end", json.dumps(analysis, ensure_ascii=False, indent=2))
                    self.txt_output.insert("end", f"\n\n📁 Word kopsavilkums saglabāts: {res_msg}\n")
                    self.txt_output.see("end")
                else:
                    self.lbl_status.configure(text="⚠️ Kļūda!")
                    self.txt_output.insert("end", f"\n❌ Kļūda failam {p.name}: {res_msg}\n")
            self.after(0, _done)

        threading.Thread(target=_worker, daemon=True).start()
