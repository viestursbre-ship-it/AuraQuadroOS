import os
import time
import threading
import json
import re
from datetime import datetime
from collections import defaultdict
from queue import Queue

# External library for .docx files
try:
    from docx import Document
    from docx.shared import Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("WARNING: python-docx library not found. Docx file processing will not work.")
    print("Please install it: pip install python-docx")
    Document = None # Placeholder to prevent errors if not installed

# External library for AI API
try:
    # Using OpenAI as a common example. Can be swapped for google.generativeai easily.
    import openai
except ImportError:
    print("WARNING: openai library not found. AI semantic analysis will not work.")
    print("Please install it: pip install openai")
    openai = None # Placeholder to prevent errors if not installed

# Define default paths
WATCHED_FOLDER = "watched_files"
DIGEST_OUTPUT_FOLDER = "storage/outbox" # Updated path as per new requirement

class EventEngine:
    """
    Centrālais notikumu dzinējs asinhronai komunikācijai starp moduļiem.
    """
    def __init__(self):
        self.subscribers = defaultdict(list)
        self.event_queue = Queue()
        self._running = False
        self._thread = None
        print("EventEngine: Sāku darbu. Gatavs notikumiem!")

    def subscribe(self, event_type, callback):
        """Reģistrē atzvanīšanas funkciju noteiktam notikuma tipam."""
        self.subscribers[event_type].append(callback)
        print(f"EventEngine: Jauns abonents {callback.__name__} notikumam '{event_type}'.")

    def publish(self, event_type, data=None):
        """Publicē notikumu, ievietojot to rindā apstrādei."""
        print(f"EventEngine: Publicēju notikumu '{event_type}' ar datiem: {data}")
        self.event_queue.put((event_type, data))

    def _process_events(self):
        """Apstrādā notikumus no rindas."""
        while self._running:
            try:
                event_type, data = self.event_queue.get(timeout=1) # Wait for 1 second
                print(f"EventEngine: Apstrādāju notikumu '{event_type}'.")
                for callback in self.subscribers[event_type]:
                    try:
                        callback(event_type, data)
                    except Exception as e:
                        print(f"EventEngine ERROR: Kļūda izpildot '{callback.__name__}' par '{event_type}': {e}")
                self.event_queue.task_done()
            except Exception as e:
                # No events in queue within timeout, or other minor issues. Keep running.
                pass

    def start(self):
        """Startē notikumu apstrādes pavedienu."""
        if not self._running:
            self._running = True
            self._thread = threading.Thread(target=self._process_events, daemon=True)
            self._thread.start()
            print("EventEngine: Notikumu apstrādes pavediens palaists.")

    def stop(self):
        """Aptur notikumu apstrādes pavedienu."""
        self._running = False
        if self._thread:
            self._thread.join(timeout=5) # Give it some time to finish current tasks
            if self._thread.is_alive():
                print("EventEngine WARNING: Pavediens neapstājās laicīgi.")
        print("EventEngine: Darbs pabeigts.")


class ServiceRegistry:
    """
    Pakalpojumu reģistrs, kas ļauj moduļiem reģistrēt un atrast citus pakalpojumus.
    """
    def __init__(self):
        self.services = {}
        print("ServiceRegistry: Sāku darbu. Gatavs reģistrēt pakalpojumus.")

    def register(self, name, service_instance):
        """Reģistrē pakalpojuma instanci ar noteiktu nosaukumu."""
        self.services[name] = service_instance
        print(f"ServiceRegistry: Pakalpojums '{name}' reģistrēts.")

    def get(self, name):
        """Atgriež reģistrēto pakalpojuma instanci pēc nosaukuma."""
        service = self.services.get(name)
        if not service:
            print(f"ServiceRegistry WARNING: Pakalpojums '{name}' nav atrasts.")
        return service
        
    def list_services(self):
        """Atgriež visu reģistrēto pakalpojumu nosaukumu sarakstu."""
        return list(self.services.keys())


class GatewayRouter:
    """
    Vienkāršs vārtejas maršrutētājs, kas apstrādā ienākošos pieprasījumus.
    Monolīta fāzē tas var apstrādāt tiešus API izsaukumus vai virzīt tos uz iekšējiem servisiem.
    """
    def __init__(self, service_registry, event_engine):
        self.service_registry = service_registry
        self.event_engine = event_engine
        print("GatewayRouter: Sāku darbu. Gatavs maršrutēt pieprasījumus.")

    def handle_request(self, path, method="GET", body=None):
        """Apstrādā ienākošo pieprasījumu un maršrutē to."""
        print(f"GatewayRouter: Saņemts pieprasījums: {method} {path}")

        # Simple routing example
        if path == "/status":
            return {"status": "OK", "services": self.service_registry.list_services()}
        elif path.startswith("/process/file"):
            if method == "POST" and body and "filepath" in body:
                # Direct call for immediate processing, or publish an event
                doc_digest_service = self.service_registry.get("DocDigestService")
                if doc_digest_service:
                    print(f"GatewayRouter: Virzu '{body['filepath']}' uz DocDigestService.")
                    # In a real scenario, this might trigger an event to decouple
                    # self.event_engine.publish("file.process", {"filepath": body["filepath"]})
                    # For a modular monolith, direct call is acceptable for synchronous ops
                    return doc_digest_service.digest_document(body["filepath"])
                else:
                    return {"error": "DocDigestService nav pieejams"}, 503
            else:
                return {"error": "Nepareizs pieprasījums /process/file"}, 400
        else:
            return {"error": "Nav atrasts"}, 404


class FolderWatcherService:
    """
    Novēro norādīto mapi jauniem failiem un publicē notikumus.
    """
    def __init__(self, watched_folder, event_engine):
        self.watched_folder = watched_folder
        self.event_engine = event_engine
        self._seen_files = set()
        self._running = False
        self._thread = None
        
        # Ensure watched folder exists
        os.makedirs(self.watched_folder, exist_ok=True)
        print(f"FolderWatcherService: Novēroju mapi '{self.watched_folder}'.")

        # Initialize with existing files to avoid reprocessing on startup
        for filename in os.listdir(self.watched_folder):
            filepath = os.path.join(self.watched_folder, filename)
            # NEW: Ignore already digested files starting with DIGEST_ on startup
            if filename.startswith("DIGEST_"):
                continue
            if os.path.isfile(filepath):
                self._seen_files.add(filepath)
        print(f"FolderWatcherService: Sākotnēji atrasti {len(self._seen_files)} faili.")

    def _watch(self):
        """Galvenā novērošanas cilpa."""
        while self._running:
            current_files = set()
            try:
                for filename in os.listdir(self.watched_folder):
                    filepath = os.path.join(self.watched_folder, filename)
                    if os.path.isfile(filepath):
                        current_files.add(filepath)

                        # CRITICAL FIX: Ignore DIGEST_ files to prevent recursion
                        if filename.startswith("DIGEST_"):
                            # print(f"FolderWatcherService: Ignorēju '{filename}', jo tas ir DIGEST_ fails.")
                            continue

                        if filepath not in self._seen_files:
                            print(f"FolderWatcherService: Atrasts jauns fails: '{filepath}'")
                            self._seen_files.add(filepath)
                            self.event_engine.publish("file.new", {"filepath": filepath})
            except FileNotFoundError:
                print(f"FolderWatcherService ERROR: Mape '{self.watched_folder}' netika atrasta.")
            except Exception as e:
                print(f"FolderWatcherService ERROR: Neparedzēta kļūda novērošanas laikā: {e}")

            # Remove deleted files from _seen_files (optional, for robustness)
            self._seen_files = {f for f in self._seen_files if os.path.exists(f)}
            
            time.sleep(5) # Check every 5 seconds

    def start(self):
        """Startē mapes novērošanas pavedienu."""
        if not self._running:
            self._running = True
            self._thread = threading.Thread(target=self._watch, daemon=True)
            self._thread.start()
            print("FolderWatcherService: Novērošanas pavediens palaists.")

    def stop(self):
        """Aptur mapes novērošanas pavedienu."""
        self._running = False
        if self._thread:
            self._thread.join(timeout=5)
        print("FolderWatcherService: Darbs pabeigts.")


class DocDigestService:
    """
    Pakalpojums, kas apstrādā dokumentus, veic AI semantisko analīzi un ģenerē strukturētus kopsavilkumus.
    """
    def __init__(self, event_engine, output_folder):
        self.event_engine = event_engine
        self.output_folder = output_folder
        # Ensure output folder exists
        os.makedirs(self.output_folder, exist_ok=True)
        self.event_engine.subscribe("file.new", self._on_new_file)
        print(f"DocDigestService: Sāku darbu. Saglabāšu kopsavilkumus '{self.output_folder}'.")

        if openai is None:
            print("DocDigestService WARNING: OpenAI bibliotēka nav pieejama, AI analīze nedarbosies.")
        if Document is None:
            print("DocDigestService WARNING: python-docx bibliotēka nav pieejama, DOCX apstrāde/ģenerēšana nedarbosies.")


    def _on_new_file(self, event_type, data):
        """Atzvanīšanas funkcija jauniem failu notikumiem."""
        filepath = data.get("filepath")
        if filepath:
            print(f"DocDigestService: Saņēmu jauna faila notikumu: '{filepath}'")
            self.digest_document(filepath)

    def _read_docx_content(self, filepath):
        """Nolasa saturu no .docx faila, ieskaitot rindkopas un tabulas."""
        if Document is None:
            print("DocDigestService ERROR: python-docx nav instalēts, nevar nolasīt .docx failu.")
            return "" 
        try:
            document = Document(filepath)
            full_text = []

            # Read paragraphs
            for para in document.paragraphs:
                full_text.append(para.text)

            # Read table cells
            for table in document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        full_text.append(cell.text)
            
            return "\n".join(full_text)
        except Exception as e:
            print(f"DocDigestService ERROR: Kļūda lasot DOCX failu '{filepath}': {e}")
            return ""

    def _read_txt_content(self, filepath):
        """Nolasa saturu no .txt faila."""
        try:
            with open(filepath, 'r', encoding='utf-8') as f:
                return f.read()
        except Exception as e:
            print(f"DocDigestService ERROR: Kļūda lasot TXT failu '{filepath}': {e}")
            return ""

    def _analyze_with_ai(self, text):
        """Veic semantisko analīzi, izmantojot LLM API (OpenAI/Gemini)."""
        if openai is None:
            return {"error": "OpenAI bibliotēka nav pieejama AI analīzei. Lūdzu, instalējiet to (pip install openai)."}

        api_key = os.getenv("AQ_AI_API_KEY")
        if not api_key:
            print("DocDigestService ERROR: Vides mainīgais 'AQ_AI_API_KEY' nav iestatīts. AI analīze nedarbosies.")
            return {"error": "AI API atslēga nav iestatīta. Lūdzu, iestatiet AQ_AI_API_KEY."}

        openai.api_key = api_key

        prompt = f"""
        Jūs esat pieredzējis juridisko dokumentu analītiķis. Izvelciet galveno informāciju no sekojošā dokumenta teksta un sniedziet to strukturētā JSON formātā.
        
        JSON struktūrai jābūt šādai:
        {{
          "document_type": "Līgums", // Piemēri: "Līgums", "Rēķins", "Vienošanās", "Cits"
          "document_number": "LV2023-001", // Ja atrasts dokumenta numurs. Ja nav, null.
          "document_date": "YYYY-MM-DD", // Ja atrasts dokumenta datums. Ja nav, null. Piemērs: "2023-10-26"
          "parties": [ // Saraksts ar pusēm un to lomām. Ja nav, tukšs masīvs.
            {{"name": "Uzņēmums A", "role": "Pakalpojumu sniedzējs"}},
            {{"name": "Uzņēmums B", "role": "Klients"}}
          ],
          "subject": "Pakalpojuma sniegšana", // Darījuma priekšmets. Ja nav, null.
          "financial_terms": {{ // Finanšu noteikumi. Ja nav, atbilstošais lauks null.
            "total_amount": "1000 EUR", // Kopējā summa (ar valūtu).
            "payment_schedule": "30 dienu laikā pēc rēķina saņemšanas", // Maksājumu grafiks vai noteikumi.
            "deadlines": "2024-12-31 (projekta pabeigšana)", // Svarīgi termiņi.
            "penalty_for_delay": "0.5% par katru kavēto dienu" // Kavējuma sankcijas.
          }},
          "ownership_termination_clauses": [ // Īpašuma tiesību un līguma izbeigšanas nosacījumi. Ja nav, tukšs masīvs.
            "Visi inteliģentā īpašuma tiesības pieder Pakalpojumu sniedzējam",
            "Līgumu var izbeigt ar 30 dienu brīdinājumu"
          ],
          "key_risks_warnings": [ // Būtiskākie riski vai brīdinājumi. Ja nav, tukšs masīvs.
            "Nav skaidri definēta atbildība par datu drošību",
            "Iespējamas izmaksas par papildu darbiem"
          ],
          "summary": "Īss, bet kodolīgs kopsavilkums par līguma galvenajiem punktiem." // Kopsavilkums.
        }}
        
        Ja lauks nav atrasts, izmantojiet `null` vai tukšu masīvu.
        Ja dokuments skaidri nav 'Līgums', 'Rēķins' vai 'Vienošanās', apzīmējiet to kā 'Cits'.
        Atbildei jābūt tikai JSON objektam.

        Dokumenta Teksts:
        ---
        {text}
        ---
        """
        try:
            # Using gpt-3.5-turbo as a cost-effective option
            response = openai.ChatCompletion.create(
                model="gpt-3.5-turbo-1106", # Updated model for JSON support
                messages=[
                    {"role": "system", "content": "You are a helpful assistant designed to output JSON."},
                    {"role": "user", "content": prompt}
                ],
                response_format={"type": "json_object"} # New feature in OpenAI API for strict JSON output
            )
            
            response_content = response.choices[0].message.content
            print(f"DocDigestService: Saņēmu AI atbildi.")
            return json.loads(response_content)
        except openai.error.OpenAIError as e:
            print(f"DocDigestService ERROR: Kļūda sazinoties ar OpenAI API: {e}")
            return {"error": f"AI API kļūda: {e}"}
        except json.JSONDecodeError as e:
            print(f"DocDigestService ERROR: Kļūda parsējot AI atbildi kā JSON: {e}\nRaw response: {response_content}")
            return {"error": "AI atbilde nav derīgs JSON."}
        except Exception as e:
            print(f"DocDigestService ERROR: Neparedzēta kļūda AI analīzes laikā: {e}")
            return {"error": f"Neparedzēta kļūda AI analīzē: {e}"}

    def _save_docx_report(self, data, original_filename_base):
        """Saglabā strukturētu atskaiti kā .docx failu ar formatējumu."""
        if Document is None:
            print("DocDigestService WARNING: python-docx bibliotēka nav pieejama, nevar saglabāt DOCX atskaiti.")
            return {"status": "error", "message": "python-docx nav instalēts."}

        document = Document()

        # Add Title
        document.add_heading(f"Dokumenta Kopsavilkums: {original_filename_base}", level=0) # level 0 is usually main title
        document.add_paragraph(f"Ģenerēts: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", style='Intense Quote')
        document.add_paragraph("\n") # Spacer

        # Document Details
        document.add_heading("1. Dokumenta Detaļas", level=1)
        doc_type = data.get("document_type", "Nav norādīts")
        doc_num = data.get("document_number", "Nav norādīts")
        doc_date = data.get("document_date", "Nav norādīts")
        document.add_paragraph(f"Tips: {doc_type}")
        document.add_paragraph(f"Numurs: {doc_num}")
        document.add_paragraph(f"Datums: {doc_date}")
        
        document.add_paragraph("\n")

        # Parties
        document.add_heading("2. Puses un to lomas", level=1)
        parties = data.get("parties", [])
        if parties:
            for party in parties:
                document.add_paragraph(f"- {party.get('name', 'Nezināma puse')} ({party.get('role', 'Nav norādīts')})", style='List Bullet')
        else:
            document.add_paragraph("Puses nav identificētas.")
        
        document.add_paragraph("\n")

        # Subject
        document.add_heading("3. Darījuma Priekšmets", level=1)
        subject = data.get("subject", "Nav norādīts")
        document.add_paragraph(subject)

        document.add_paragraph("\n")

        # Financial Terms
        document.add_heading("4. Finanšu Noteikumi", level=1)
        financial_terms = data.get("financial_terms", {})
        total_amount = financial_terms.get("total_amount", "Nav norādīts")
        payment_schedule = financial_terms.get("payment_schedule", "Nav norādīts")
        deadlines = financial_terms.get("deadlines", "Nav norādīts")
        penalty = financial_terms.get("penalty_for_delay", "Nav norādīts")

        document.add_paragraph(f"- Kopējā summa: {total_amount}", style='List Bullet')
        document.add_paragraph(f"- Maksājumu grafiks: {payment_schedule}", style='List Bullet')
        document.add_paragraph(f"- Termiņi: {deadlines}", style='List Bullet')
        document.add_paragraph(f"- Kavējuma sankcijas: {penalty}", style='List Bullet')

        document.add_paragraph("\n")

        # Ownership & Termination
        document.add_heading("5. Īpašuma Tiesības un Līguma Izbeigšana", level=1)
        ownership_termination = data.get("ownership_termination_clauses", [])
        if ownership_termination:
            for clause in ownership_termination:
                document.add_paragraph(f"- {clause}", style='List Bullet')
        else:
            document.add_paragraph("Nav īpašu noteikumu.")

        document.add_paragraph("\n")

        # Risks & Warnings
        document.add_heading("6. Būtiskākie Riski vai Brīdinājumi", level=1)
        risks_warnings = data.get("key_risks_warnings", [])
        if risks_warnings:
            for risk in risks_warnings:
                document.add_paragraph(f"- {risk}", style='List Bullet')
        else:
            document.add_paragraph("Nav būtisku risku vai brīdinājumu.")
        
        document.add_paragraph("\n")

        # Summary
        document.add_heading("7. Kopsavilkums", level=1)
        summary = data.get("summary", "Kopsavilkums nav pieejams.")
        document.add_paragraph(summary)

        # Save the document
        output_filename = f"DIGEST_{original_filename_base}.docx"
        output_filepath = os.path.join(self.output_folder, output_filename)
        try:
            document.save(output_filepath)
            print(f"DocDigestService: Saglabāju kopsavilkumu kā '{output_filepath}'.")
            return {"status": "success", "output_filepath": output_filepath}
        except Exception as e:
            print(f"DocDigestService ERROR: Kļūda saglabājot DOCX atskaiti '{output_filepath}': {e}")
            return {"status": "error", "message": f"Kļūda saglabājot DOCX atskaiti: {e}"}


    def digest_document(self, filepath):
        """Galvenā funkcija dokumenta apstrādei: lasīšana, AI analīze, atskaites ģenerēšana."""
        file_extension = os.path.splitext(filepath)[1].lower()
        original_filename_base = os.path.splitext(os.path.basename(filepath))[0]
        document_text = ""

        if file_extension == ".docx":
            document_text = self._read_docx_content(filepath)
        elif file_extension == ".txt":
            document_text = self._read_txt_content(filepath)
        else:
            print(f"DocDigestService WARNING: neatbalstīts faila tips: '{filepath}'. Apstrādāts tikai kā teksta fails.")
            document_text = self._read_txt_content(filepath) # Fallback to txt read

        if not document_text.strip():
            print(f"DocDigestService: Fails '{filepath}' ir tukšs vai neizdevās nolasīt.")
            return {"status": "error", "message": "Fails ir tukšs vai neizdevās nolasīt."}

        # Perform AI analysis
        ai_analysis_result = self._analyze_with_ai(document_text)

        if ai_analysis_result.get("error"):
            return ai_analysis_result # Pass AI error directly

        # Save as DOCX report
        report_result = self._save_docx_report(ai_analysis_result, original_filename_base)
        
        self.event_engine.publish("document.digested", {
            "original_filepath": filepath,
            "digest_output": report_result.get("output_filepath"),
            "summary": ai_analysis_result.get("summary", "Nav kopsavilkuma."),
            "type": ai_analysis_result.get("document_type", "Cits")
        })
        return report_result


class AQOSModularMonolith:
    """
    AQ-OS Modulārā Monolīta galvenā lietojumprogramma, kas apvieno visus pakalpojumus.
    """
    def __init__(self):
        print("AQ-OS Modular Monolith: Iniciēju sistēmas komponentes...")
        self.event_engine = EventEngine()
        self.service_registry = ServiceRegistry()
        self.gateway_router = GatewayRouter(self.service_registry, self.event_engine)

        # Register core services
        self.service_registry.register("EventEngine", self.event_engine)
        self.service_registry.register("ServiceRegistry", self.service_registry)
        self.service_registry.register("GatewayRouter", self.gateway_router)

        # Initialize and register FolderWatcherService
        self.folder_watcher = FolderWatcherService(WATCHED_FOLDER, self.event_engine)
        self.service_registry.register("FolderWatcherService", self.folder_watcher)

        # Initialize and register DocDigestService
        self.doc_digest_service = DocDigestService(self.event_engine, DIGEST_OUTPUT_FOLDER)
        self.service_registry.register("DocDigestService", self.doc_digest_service)

    def start(self):
        """Startē visas monolīta komponentes."""
        print("AQ-OS Modular Monolith: Startēju sistēmu...")
        self.event_engine.start()
        self.folder_watcher.start()
        print("AQ-OS Modular Monolith: Sistēma palaista! Gaidu failus mapē 'watched_files'.")
        print(f"Kopsavilkumi tiks saglabāti mapē '{DIGEST_OUTPUT_FOLDER}'.")
        print("\n*** SVARĪGI: Lai AI analīze darbotos, iestatiet vides mainīgo AQ_AI_API_KEY! ***")
        print("Piemērs: export AQ_AI_API_KEY='your-openai-api-key'\n")
        
        # Keep the main thread alive, or use a simple loop
        try:
            while True:
                time.sleep(1)
        except KeyboardInterrupt:
            print("\nAQ-OS Modular Monolith: Saņemts apturēšanas pieprasījums.")
            self.stop()

    def stop(self):
        """Aptur visas monolīta komponentes."""
        print("AQ-OS Modular Monolith: Apturu sistēmas komponentes...")
        self.folder_watcher.stop()
        self.event_engine.stop()
        print("AQ-OS Modular Monolith: Sistēma apturēta. Arrivederci!")

# Entry point
if __name__ == "__main__":
    monolith = AQOSModularMonolith()
    monolith.start()