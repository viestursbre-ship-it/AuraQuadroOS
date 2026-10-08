import os
import sys
import time
import json
import shutil
import threading
from pathlib import Path
from datetime import datetime, timedelta

# GUI un sistēmas joslas bibliotēkas
import customtkinter as ctk
from PIL import Image, ImageDraw
import pystray
from pystray import MenuItem as item

# Biroja un dokumentu bibliotēkas
import docx
from docx.shared import Inches, Pt
from pptx import Presentation
import openpyxl
from openpyxl.styles import Font

# MI un Windows COM integrācija
from google import genai
from google.genai import types

try:
    import win32com.client
    import pythoncom
    HAS_WIN32COM = True
except ImportError:
    HAS_WIN32COM = False

# ==========================================
# KONFIGURĀCIJA UN CEĻI
# ==========================================
BASE_DIR = Path(__file__).resolve().parent
STORAGE_DIR = BASE_DIR / "storage"
INBOX_DIR = STORAGE_DIR / "inbox"
OUTBOX_DIR = STORAGE_DIR / "outbox"
ARCHIVE_DIR = STORAGE_DIR / "archive"
EXCEL_REGISTRY_PATH = STORAGE_DIR / "Ligumu_Registrs.xlsx"
STATE_FILE = BASE_DIR / "radar_state.json"
IGNORE_FILE = BASE_DIR / "radar_ignored_senders.json"

for d in [INBOX_DIR, OUTBOX_DIR, ARCHIVE_DIR]:
    d.mkdir(parents=True, exist_ok=True)

# ==========================================
# IGNORĒTO SŪTĪTĀJU PĀRVALDĪBA
# ==========================================
def get_ignored_senders():
    if not os.path.exists(IGNORE_FILE):
        return set()
    try:
        with open(IGNORE_FILE, "r", encoding="utf-8") as f:
            return set(json.load(f))
    except Exception:
        return set()

def add_ignored_sender(sender_identifier):
    if not sender_identifier:
        return
    ignored = get_ignored_senders()
    ignored.add(sender_identifier.strip().lower())
    with open(IGNORE_FILE, "w", encoding="utf-8") as f:
        json.dump(list(ignored), f, ensure_ascii=False, indent=2)

# ==========================================
# 1. DOKUMENTU APSTRĀDES MODULIS
# ==========================================
def safe_archive_file(source_path: Path, target_dir: Path) -> Path:
    target_dir.mkdir(parents=True, exist_ok=True)
    destination = target_dir / source_path.name
    if destination.exists():
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        destination = target_dir / f"{source_path.stem}_{timestamp}{source_path.suffix}"
    shutil.move(str(source_path), str(destination))
    return destination

def append_to_excel_registry(file_name: str, analysis: dict, status: str = "Apstrādāts"):
    headers = ["Ieraksta Datums", "Fails", "Dokumenta tips", "Dokumenta Nr.", 
               "Dokumenta Datums", "Puses", "Summa", "Termiņš", "Statuss"]
    try:
        if not EXCEL_REGISTRY_PATH.exists():
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Līgumu Reģistrs"
            ws.append(headers)
            for cell in ws[1]:
                cell.font = Font(bold=True)
            wb.save(EXCEL_REGISTRY_PATH)

        wb = openpyxl.load_workbook(EXCEL_REGISTRY_PATH)
        ws = wb.active

        parties_data = analysis.get("parties", [])
        if isinstance(parties_data, list):
            parties_str = ", ".join([p.get("name", "").strip() if isinstance(p, dict) else str(p) for p in parties_data])
        else:
            parties_str = str(parties_data or "N/A")

        fin = analysis.get("financial_terms", {})
        sum_val = fin.get("total_amount", "N/A") if isinstance(fin, dict) else "N/A"
        term_val = fin.get("deadlines", "N/A") if isinstance(fin, dict) else "N/A"

        row = [
            datetime.now().strftime("%Y-%m-%d %H:%M"),
            file_name,
            str(analysis.get("document_type", "N/A")),
            str(analysis.get("document_number", "N/A")),
            str(analysis.get("document_date", "N/A")),
            parties_str,
            str(sum_val),
            str(term_val),
            status
        ]
        ws.append(row)
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)
        wb.save(EXCEL_REGISTRY_PATH)
    except Exception as e:
        print(f"⚠️ Kļūda Excel reģistrācijā: {e}")

class DocDigestService:
    def _read_docx(self, path: Path) -> str:
        doc = docx.Document(path)
        full_text = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Nolasām arī tabulu saturu (divvalodu līgumiem)
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    full_text.append(" | ".join(cells))
        return "\n".join(full_text)

    def _read_txt(self, path: Path) -> str:
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    def _read_pptx(self, path: Path) -> str:
        prs = Presentation(path)
        full_text = []
        for i, slide in enumerate(prs.slides, 1):
            full_text.append(f"\n--- Slaids {i} ---")
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        if paragraph.text.strip():
                            full_text.append(paragraph.text.strip())
        return "\n".join(full_text)

    def _analyze_with_gemini(self, content_payload, file_type: str = "docx") -> dict:
        api_key = os.getenv("AQ_AI_API_KEY")
        if not api_key:
            return {"error": "AQ_AI_API_KEY mainīgais nav atrasts sistēmā."}

        client = genai.Client(api_key=api_key)
        prompt_text = """Analizējiet šo dokumentu un sagatavojiet precīzu JSON objektu:
{
  "document_type": "Līgums / Rēķins / Akts / Cits",
  "document_number": "Numurs vai N/A",
  "document_date": "Datums vai N/A",
  "parties": [{"name": "Puses nosaukums", "role": "Pasūtītājs/Piegādātājs"}],
  "subject": "Darījuma priekšmets",
  "financial_terms": {"total_amount": "Summa ar PVN", "deadlines": "Apmaksas/Izpildes termiņi"},
  "risks_warnings": ["Konstatētie riski vai brīdinājumi"]
}
Atgrieziet TIKAI derīgu JSON bez markdown blokiem."""

        try:
            if file_type == "pdf":
                response = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=[prompt_text, content_payload],
                    config=types.GenerateContentConfig(temperature=0.2)
                )
            else:
                response = client.models.generate_content(
                    model="gemini-2.5-flash",
                    contents=f"{prompt_text}\n\nDokumenta saturs:\n{content_payload[:15000]}",
                    config=types.GenerateContentConfig(temperature=0.2)
                )
            if response and response.text:
                clean_json = response.text.replace("```json", "").replace("```", "").strip()
                return json.loads(clean_json)
            return {"error": "Tukša atbilde no MI"}
        except Exception as e:
            return {"error": str(e)}

    def _create_digest_docx(self, output_path: Path, data: dict, original_filename: str):
        doc = docx.Document()
        doc.add_heading(f"Kopsavilkums: {original_filename}", level=1)
        doc.add_paragraph(f"Izveidots: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")

        def add_sec(title, val):
            doc.add_heading(title, level=2)
            if isinstance(val, dict):
                for k, v in val.items():
                    doc.add_paragraph(f"• {k.replace('_', ' ').capitalize()}: {v}")
            elif isinstance(val, list):
                for item in val:
                    doc.add_paragraph(f"• {item}")
            else:
                doc.add_paragraph(str(val or "N/A"))

        add_sec("Vispārīgā informācija", {
            "Tips": data.get("document_type", "N/A"),
            "Numurs": data.get("document_number", "N/A"),
            "Datums": data.get("document_date", "N/A"),
            "Priekšmets": data.get("subject", "N/A")
        })
        add_sec("Finanšu noteikumi", data.get("financial_terms", {}))
        add_sec("Riski un brīdinājumi", data.get("risks_warnings", []))
        doc.save(output_path)

    def process_file(self, file_path: Path):
        ext = file_path.suffix.lower()
        text_payload = None

        print(f"📄 [DOKUMENTI] Sāk analīzi: {file_path.name}")
        if ext == ".docx":
            text_payload = self._read_docx(file_path)
        elif ext == ".txt":
            text_payload = self._read_txt(file_path)
        elif ext == ".pptx":
            text_payload = self._read_pptx(file_path)
        elif ext == ".pdf":
            with open(file_path, "rb") as f:
                pdf_bytes = f.read()
            text_payload = types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf")
        else:
            return

        if not text_payload:
            print(f"⚠️ [DOKUMENTI] Neizdevās iegūt tekstu no {file_path.name}")
            return

        analysis = self._analyze_with_gemini(text_payload, file_type=ext.replace(".", ""))
        if "error" in analysis:
            print(f"⚠️ [MI KĻŪDA] {file_path.name}: {analysis['error']}")
            return

        out_name = f"DIGEST_{file_path.stem}.docx"
        self._create_digest_docx(OUTBOX_DIR / out_name, analysis, file_path.name)
        append_to_excel_registry(file_path.name, analysis)
        safe_archive_file(file_path, ARCHIVE_DIR)
        print(f"✅ [DOKUMENTI] Pabeigts: {file_path.name} -> {out_name}")

class FolderWatcherService:
    def __init__(self, watch_dir: Path, digest_service: DocDigestService):
        self.watch_dir = watch_dir
        self.digest_service = digest_service
        self.running = False

    def start_watching(self, interval_sec: int = 5):
        self.running = True
        print(f"👀 [DOKUMENTI] Mape {self.watch_dir} tiek uzraudzīta...")
        while self.running:
            try:
                if self.watch_dir.exists():
                    for file_path in list(self.watch_dir.iterdir()):
                        if file_path.is_file() and not file_path.name.startswith("~"):
                            try:
                                self.digest_service.process_file(file_path)
                            except Exception as file_err:
                                print(f"⚠️ Kļūda failam {file_path.name}: {file_err}")
            except Exception as e:
                print(f"⚠️ Uzrauga kļūda: {e}")
            time.sleep(interval_sec)

# ==========================================
# 2. OUTLOOK RADARA SKENERIS (UNIVERSĀLS & VISAS MAPES)
# ==========================================
def clean_msg_text(msg):
    """Izvelk un sakārto tekstu, noņemot kodējuma ķeburus."""
    text = ""
    try:
        if hasattr(msg, "Body") and msg.Body:
            text = msg.Body
        elif hasattr(msg, "HTMLBody") and msg.HTMLBody:
            import re
            text = re.sub(r'<[^<]+?>', '', msg.HTMLBody)
    except Exception:
        text = ""
    
    # Salabojam Windows/MAPI latviešu valodas aizvietojumus
    text = text.replace(r"\d", "ī").replace(r"\t", "ī").replace(r"\i", "ī").replace(r"\j", "ā")
    return text.strip()

def collect_all_folders(root_folder):
    """Rekursīvi atrod un atgriež mapi un pilnīgi visas tās apakšmapes."""
    folders = [root_folder]
    try:
        for sub in root_folder.Folders:
            folders.extend(collect_all_folders(sub))
    except Exception:
        pass
    return folders

def scan_outlook_mailbox():
    if not HAS_WIN32COM:
        return

    pythoncom.CoInitialize()
    try:
        outlook = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
        inbox = outlook.GetDefaultFolder(6)
        
        # Paņemam gan Inbox, gan visas apakšmapes (Support, Zebra, Sales utt.)
        all_target_folders = collect_all_folders(inbox)

        # Meklējam pēdējo 30 dienu e-pastus ar MAPI filtru
        start_date = (datetime.now() - timedelta(days=30)).strftime("%d/%m/%Y")
        filter_query = f"[ReceivedTime] >= '{start_date}'"

        items_to_save = {}
        if os.path.exists(STATE_FILE):
            try:
                with open(STATE_FILE, "r", encoding="utf-8") as f:
                    items_to_save = json.load(f)
            except Exception:
                pass

        ignored_list = get_ignored_senders()
        new_count = 0

        for folder in all_target_folders:
            folder_name = getattr(folder, "Name", "Nezināma")
            try:
                messages = folder.Items
                filtered_messages = messages.Restrict(filter_query)
                filtered_messages.Sort("[ReceivedTime]", True)
                total = filtered_messages.Count
            except Exception:
                continue

            for i in range(1, total + 1):
                try:
                    msg = filtered_messages.Item(i)
                    entry_id = msg.EntryID
                    if entry_id in items_to_save:
                        continue

                    subj = msg.Subject or "(Bez temata)"
                    sender = msg.SenderName or "Nezināms"
                    body = clean_msg_text(msg)

                    if any(ign in sender.lower() for ign in ignored_list):
                        continue

                    subj_lower = subj.lower()
                    body_lower = body.lower()
                    
                    # 1. Vai tas ir paša Viestura uzdevums no telefona?
                    is_self_task = (
                        "viesturs" in sender.lower() or 
                        "bredovskis" in sender.lower() or 
                        "outlook for android" in body_lower or
                        "līdz rītdienai" in body_lower or
                        "radar" in subj_lower
                    )

                    # 2. Standarta biznesa atslēgvārdi
                    keywords = ["cenu piepras", "iepirkums", "termin", "vid", "pasutijum", "līgums", "rekins", "rēķin"]
                    matches_keywords = any(k in subj_lower for k in keywords)

                    # Ja atbilst vai nu atslēgvārdiem, VAI tas ir tiešais sūtījums no telefona
                    if matches_keywords or is_self_task:
                        import re
                        date_match = re.search(r'\b(\d{1,2}[\./]\d{1,2}(?:[\./]\d{2,4})?)\b', subj + " " + body[:300])
                        
                        if date_match:
                            found_deadline = f"Līdz {date_match.group(1)}"
                        elif "rīt" in body_lower or "rītdien" in body_lower:
                            found_deadline = "Līdz rītdienai"
                        else:
                            found_deadline = "Steidzams"

                        cat = "Zibens uzdevums" if is_self_task else ("Cenu pieprasījums" if "cenu" in subj_lower else "Kritisks")

                        items_to_save[entry_id] = {
                            "subject": subj,
                            "sender": f"{sender} ({folder_name})",
                            "core_request": body[:120].strip() + "..." if body else subj,
                            "deadline": found_deadline,
                            "category": cat,
                            "status": "ACTIVE",
                            "received": msg.ReceivedTime.strftime("%d.%m %H:%M") if hasattr(msg, "ReceivedTime") else ""
                        }
                        new_count += 1
                        print(f"🎯 [RADARS] Atrasts [{folder_name}]: {subj[:45]}")
                except Exception:
                    continue

        if new_count > 0:
            with open(STATE_FILE, "w", encoding="utf-8") as f:
                json.dump(items_to_save, f, ensure_ascii=False, indent=2)
            print(f"💾 [RADARS] Saglabāti {new_count} jauni uzdevumi.")

    except Exception as e:
        print(f"⚠️ [RADARS KĻŪDA]: {e}")
    finally:
        pythoncom.CoUninitialize()

class OutlookRadarWatcher(threading.Thread):
    def __init__(self, interval_sec: int = 180):
        super().__init__(daemon=True)
        self.interval_sec = interval_sec

    def run(self):
        print("🛰️ [RADARS] Fona skeneris palaists...")
        while True:
            try:
                scan_outlook_mailbox()
            except Exception as e:
                print(f"⚠️ [RADARS CIKLA KĻŪDA]: {e}")
            time.sleep(self.interval_sec)
        
# ==========================================
# 3. AQ MISSION CONTROL / COCKPIT UI
# ==========================================
def open_email_in_outlook(entry_id):
    if not HAS_WIN32COM:
        return
    try:
        pythoncom.CoInitialize()
        outlook = win32com.client.Dispatch("Outlook.Application").GetNamespace("MAPI")
        msg = outlook.GetItemFromID(entry_id)
        msg.Display()
    except Exception as e:
        print(f"⚠️ Nevarēja atvērt e-pastu: {e}")
    finally:
        pythoncom.CoUninitialize()

class RadarApp(ctk.CTk):
    def __init__(self, digest_service: DocDigestService):
        super().__init__()
        self.digest_service = digest_service
        self.title("AQ-OS Cockpit v1.0")
        self.geometry("480x680")
        self.resizable(False, False)
        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        sw = self.winfo_screenwidth()
        sh = self.winfo_screenheight()
        self.geometry(f"480x680+{sw - 500}+{sh - 760}")
        self.protocol("WM_DELETE_WINDOW", self.withdraw)

        # Cockpit Augšējā josla
        header = ctk.CTkFrame(self, fg_color="#0f172a", corner_radius=0)
        header.pack(fill="x", padx=0, pady=0)

        ctk.CTkLabel(header, text="⚡ AQ-OS MISSION CONTROL", font=ctk.CTkFont(size=14, weight="bold"), text_color="#38bdf8").pack(side="left", padx=14, pady=10)
        
        # Vadības pogas
        btn_box = ctk.CTkFrame(header, fg_color="transparent")
        btn_box.pack(side="right", padx=10)
        
        ctk.CTkButton(btn_box, text="📥 Skenēt Pastu", width=95, height=26, fg_color="#2563eb", hover_color="#1d4ed8", font=ctk.CTkFont(size=11), command=self.trigger_scan).pack(side="left", padx=3)
        ctk.CTkButton(btn_box, text="🔄", width=32, height=26, fg_color="#334155", command=self.load_cards).pack(side="left", padx=2)

        # Saraksta lauks
        self.scroll_frame = ctk.CTkScrollableFrame(self, width=450, height=580)
        self.scroll_frame.pack(padx=10, pady=10, fill="both", expand=True)
        self.load_cards()

    def trigger_scan(self):
        threading.Thread(target=self._manual_scan_task, daemon=True).start()

    def _manual_scan_task(self):
        print("🛰️ [COCKPIT] Manuāla pasta sinhronizācija palaista...")
        scan_outlook_mailbox()
        self.after(500, self.load_cards)

    def show_window(self):
        self.deiconify()
        self.lift()
        self.load_cards()

    def mark_as_done(self, entry_id, card_widget):
        if os.path.exists(STATE_FILE):
            try:
                with open(STATE_FILE, "r", encoding="utf-8") as f:
                    data = json.load(f)
                if entry_id in data:
                    data[entry_id]["status"] = "DONE"
                    with open(STATE_FILE, "w", encoding="utf-8") as f:
                        json.dump(data, f, ensure_ascii=False, indent=2)
            except Exception:
                pass
        card_widget.destroy()

    def ignore_sender_and_remove(self, sender, entry_id, card_widget):
        add_ignored_sender(sender)
        self.mark_as_done(entry_id, card_widget)

    def load_cards(self):
        for w in self.scroll_frame.winfo_children():
            w.destroy()

        if not os.path.exists(STATE_FILE):
            ctk.CTkLabel(self.scroll_frame, text="Nav aktīvu datu.").pack(pady=30)
            return

        try:
            with open(STATE_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)
        except Exception:
            data = {}

        ignored_list = get_ignored_senders()
        active_items = [
            (k, v) for k, v in data.items() 
            if v.get("status") == "ACTIVE" and not any(ign in v.get("sender", "").lower() for ign in ignored_list)
        ]

        # Kārtojam: jaunākie datumi pašā augšā
        def parse_date(item_tuple):
            val = item_tuple[1].get("received", "")
            try:
                return datetime.strptime(f"2026.{val}", "%Y.%d.%m %H:%M")
            except Exception:
                return datetime.min

        active_items.sort(key=parse_date, reverse=True)

        if not active_items:
            ctk.CTkLabel(self.scroll_frame, text="✨ Visi uzdevumi nokārtoti!", font=ctk.CTkFont(size=14)).pack(pady=40)
            return

        for entry_id, item_data in active_items:
            card = ctk.CTkFrame(self.scroll_frame, corner_radius=8, fg_color="#131b2e")
            card.pack(pady=6, padx=4, fill="x")

            cat = item_data.get("category", "Cits")
            badge_bg = "#ef4444" if cat in ["Līgums", "Kritisks"] else "#f59e0b" if "pieprasījums" in cat.lower() else "#3b82f6"

            top_bar = ctk.CTkFrame(card, fg_color="transparent")
            top_bar.pack(fill="x", padx=10, pady=(8, 2))
            ctk.CTkLabel(top_bar, text=f" {cat} ", fg_color=badge_bg, corner_radius=4, text_color="white", font=ctk.CTkFont(size=10, weight="bold")).pack(side="left")

            dl = item_data.get("deadline")
            if dl:
                ctk.CTkLabel(
                    top_bar, 
                    text=f" ⏳ {dl} ", 
                    fg_color="#dc2626", 
                    corner_radius=4, 
                    text_color="white", 
                    font=ctk.CTkFont(size=10, weight="bold")
                ).pack(side="left", padx=(6, 0))

            if item_data.get("received"):
                ctk.CTkLabel(top_bar, text=item_data.get("received"), text_color="#64748b", font=ctk.CTkFont(size=10)).pack(side="right")

            ctk.CTkLabel(card, text=item_data.get("subject", ""), font=ctk.CTkFont(size=12, weight="bold"), wraplength=410, justify="left").pack(anchor="w", padx=10, pady=(2, 0))
            ctk.CTkLabel(card, text=f"No: {item_data.get('sender', '')}", font=ctk.CTkFont(size=11), text_color="#94a3b8", wraplength=410, justify="left").pack(anchor="w", padx=10)

            req = item_data.get("core_request")
            if req:
                ctk.CTkLabel(card, text=req, font=ctk.CTkFont(size=11), text_color="#cbd5e1", wraplength=410, justify="left").pack(anchor="w", padx=10, pady=4)

            btn_row = ctk.CTkFrame(card, fg_color="transparent")
            btn_row.pack(fill="x", padx=10, pady=(4, 8))
            
            ctk.CTkButton(btn_row, text="↗ Atvērt", width=75, height=24, fg_color="#2563eb", hover_color="#1d4ed8", font=ctk.CTkFont(size=11),
                          command=lambda eid=entry_id: open_email_in_outlook(eid)).pack(side="left", padx=(0, 6))
            
            ctk.CTkButton(btn_row, text="✓ Nokārtots", width=85, height=24, fg_color="#059669", hover_color="#047857", font=ctk.CTkFont(size=11),
                          command=lambda eid=entry_id, c=card: self.mark_as_done(eid, c)).pack(side="left", padx=(0, 6))

            ctk.CTkButton(btn_row, text="🚫 Ignorēt", width=75, height=24, fg_color="#334155", hover_color="#dc2626", font=ctk.CTkFont(size=11),
                          command=lambda s=item_data.get("sender", ""), eid=entry_id, c=card: self.ignore_sender_and_remove(s, eid, c)).pack(side="left")

def run_tray(app):
    img = Image.new("RGBA", (64, 64), color=(0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    d.ellipse((6, 6, 58, 58), fill="#f59e0b")

    def on_open(icon, itm):
        app.after(0, app.show_window)

    def on_quit(icon, itm):
        icon.stop()
        app.after(0, app.destroy)

    menu = pystray.Menu(item("Atvērt Cockpit", on_open, default=True), item("Iziet", on_quit))
    icon = pystray.Icon("AQ_Cockpit", img, "AQ-OS Cockpit", menu)
    icon.run()

# ==========================================
# 4. GALVENAIS STARTS
# ==========================================
if __name__ == "__main__":
    print("🚀 [AQ-OS] Mission Control startēts...")

    digest_service = DocDigestService()
    watcher = FolderWatcherService(INBOX_DIR, digest_service)
    doc_thread = threading.Thread(target=watcher.start_watching, daemon=True)
    doc_thread.start()

    radar_thread = OutlookRadarWatcher(interval_sec=300)
    radar_thread.start()

    app = RadarApp(digest_service)
    tray_thread = threading.Thread(target=run_tray, args=(app,), daemon=True)
    tray_thread.start()

    app.mainloop()
